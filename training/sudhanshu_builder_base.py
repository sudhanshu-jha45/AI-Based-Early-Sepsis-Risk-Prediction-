import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_p(doc, text, style='Normal', space_before=4, space_after=6, line_spacing=1.15, bold=False, italic=False, color_rgb=(0,0,0)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    if color_rgb != (0,0,0):
        run.font.color.rgb = RGBColor(*color_rgb)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(30, 58, 138) # Dark Blue
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return p

def add_formula_box(doc, eq_label, formula_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    
    run_lbl = cp.add_run(f"{eq_label}:  ")
    run_lbl.font.name = 'Calibri'
    run_lbl.font.bold = True
    run_lbl.font.size = Pt(10.5)
    run_lbl.font.color.rgb = RGBColor(30, 58, 138)
    
    run_eq = cp.add_run(formula_text)
    run_eq.font.name = 'Cambria Math'
    run_eq.font.italic = True
    run_eq.font.size = Pt(11)
    run_eq.font.color.rgb = RGBColor(15, 23, 42)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_code_box(doc, code_title, code_text):
    p_lbl = doc.add_paragraph()
    p_lbl.paragraph_format.space_before = Pt(8)
    p_lbl.paragraph_format.space_after = Pt(2)
    p_lbl.paragraph_format.keep_with_next = True
    r_lbl = p_lbl.add_run(code_title)
    r_lbl.font.name = 'Consolas'
    r_lbl.font.bold = True
    r_lbl.font.size = Pt(10)
    r_lbl.font.color.rgb = RGBColor(30, 58, 138)
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.line_spacing = 1.05
    run = cp.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_custom_table(doc, headers, rows_data):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.name = 'Calibri'
        
    for r_idx, row_vals in enumerate(rows_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.size = Pt(9)
                p.runs[0].font.name = 'Calibri'
            
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table

print("Sudhanshu docx builder base initialized successfully.")
