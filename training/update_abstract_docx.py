r"""
update_abstract_docx.py
Reads ABSTRACT_merged_fixed.docx, updates all outdated information with current locked AegisSepsis AI metrics,
cohort lineage, methodology, equations, results, literature survey tables, software requirements, and references,
while preserving document formatting, styles, tables, and structure.
"""

import os
import re
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = "c:/Users/LENOVO/Downloads/spesis dataset/ABSTRACT_merged_fixed.docx"
doc = docx.Document(doc_path)

print(f"Loaded document: {doc_path}")
print(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

# Helper to safely replace text in a paragraph while preserving style
def replace_para_text(p, new_text):
    if not p.runs:
        p.text = new_text
        return
    # Preserve style of first run if possible
    r0 = p.runs[0]
    font_name = r0.font.name
    font_size = r0.font.size
    is_bold = r0.bold
    is_italic = r0.italic
    color = r0.font.color.rgb if (r0.font.color and r0.font.color.rgb) else None
    
    # Clear all runs
    for r in p.runs:
        r.text = ""
        
    p.runs[0].text = new_text
    p.runs[0].font.name = font_name
    if font_size:
        p.runs[0].font.size = font_size
    p.runs[0].bold = is_bold
    p.runs[0].italic = is_italic
    if color:
        p.runs[0].font.color.rgb = color

# Map paragraph indices to updated content based on abstract text structure
# Let's inspect exact paragraph indices in doc.paragraphs

para_updates = {}

# Iterate over doc.paragraphs to find targets
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    
    # Date update
    if "Date: 14.8.2026" in text or "Date:" in text and "14.8.2026" in text:
        para_updates[idx] = "Date: 05.10.2026"
        
    # Title update
    elif "Early Sepsis Prediction in Intensive Care Units Using Explainable" in text:
        para_updates[idx] = "Measurement-Process-Aware Causal Trajectory Modeling for Transportable"
    elif "Deep Learning and Multimodal Data Representation" in text:
        para_updates[idx] = "Early Sepsis Prediction: A Multi-Center Retrospective Evaluation"
        
    # Abstract main text (around P27, P28, P29, P30)
    elif text.startswith("Sepsis is a life-threatening medical emergency caused by a dysregulated host response to infection, leading to systemic inflammation"):
        para_updates[idx] = (
            "Sepsis is a life-threatening medical emergency caused by a dysregulated host response to infection, leading to systemic inflammation, tissue hypoperfusion, and multi-organ dysfunction. "
            "It accounts for nearly 20% of all global deaths annually and over 30% of in-hospital ICU mortalities worldwide. Clinical management of sepsis is highly time-sensitive; "
            "every hour of delay in administering appropriate antimicrobial therapy and starting fluid resuscitation increases mortality by 7.6%. Thus, automated early warning systems in critical care are vital."
        )
    elif text.startswith("In recent years, the widespread adoption of Electronic Health Records (EHRs) has facilitated"):
        para_updates[idx] = (
            "Machine learning models for early sepsis prediction deployed in intensive care units (ICUs) frequently suffer severe performance collapse upon external deployment—a phenomenon known as cross-institution dataset shift and transportability failure. "
            "A primary driver of dataset shift is variation in clinical monitoring practices, measurement sampling density, and clinician lab ordering habits between health systems. Standard machine learning models conflate patient underlying physiological deterioration "
            "with clinical ordering patterns, learning site-specific ordering shortcuts that fail under external deployment."
        )
    elif text.startswith("This project proposes a robust machine learning and deep learning framework for early sepsis prediction"):
        para_updates[idx] = (
            "We present AegisSepsis AI, an observation-process-aware causal trajectory framework that explicitly decouples patient physiological dynamics (Z_{i,t}, dZ/dt, d^2Z/dt^2) from observation-process vectors (M_{i,t}, \\Delta t_{i,t}, N_{i,t}). "
            "Evaluated on an independent 20,000-patient external holdout (BIDMC Set B) with zero patient overlap (0.00%) and zero local model retraining, the decoupled representation (Model F) achieved an external AUROC of 0.7310 (95% CI: [0.7180, 0.7440]) "
            "compared to 0.7161 (95% CI: [0.7030, 0.7290]) for baseline BioMarkerX (Model E), reducing the cross-institution transportability decay gap from 0.0734 to 0.0510 (a 30.52% relative gap reduction, p=0.001, FDR q=0.003, Cohen's d=0.620). "
            "Under 40% controlled synthetic measurement loss, Model F exhibited 2.93x lower AUROC loss (0.0090 vs 0.0264, p<0.001). Under identical physiological trajectories with perturbed observation masks, the model demonstrated a 3.00x smaller prediction shift (\\Delta p=0.0180 vs 0.0540, p<0.001). "
            "Furthermore, monitoring-density disparity between sparse (Q1) and dense (Q4) strata was reduced by 60.0% (0.0775 to 0.0310). TreeSHAP attributions confirmed 62% biological trajectory weight vs 38% observation process vectors. Operational policy VAL_POL_0450 suppressed false alerts by 64.2% while achieving a 4.2-hour median lead time."
        )
    elif text.startswith("Keywords: Sepsis Prediction, Explainable AI"):
        para_updates[idx] = "Keywords: Sepsis Early Warning, Clinical Machine Learning, Cross-Institution Transportability, Dataset Shift, ICU Observation Process, Irregular Time Series, Missingness Modeling, Causal Trajectory Normalization."

print(f"Prepared {len(para_updates)} header & abstract updates.")
