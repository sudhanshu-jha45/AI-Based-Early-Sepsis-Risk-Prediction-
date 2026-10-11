import docx

doc = docx.Document(r"c:\Users\LENOVO\Downloads\spesis dataset\REVIEW-2 DOCUMENT FORMAT-Corrected.docx")

print("=== ALL PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    print(f"[{i}] ({p.style.name}): {repr(p.text)}")

print("\n=== ALL TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows):
        row_text = [cell.text.replace('\n', ' ') for cell in row.cells]
        print(f"Row {r_idx}: {row_text}")
