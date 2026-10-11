import docx

doc = docx.Document(r"c:\Users\LENOVO\Downloads\spesis dataset\REVIEW-2 DOCUMENT FORMAT-Corrected.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P[{i}] (Style: {p.style.name}): {p.text[:100]}")
