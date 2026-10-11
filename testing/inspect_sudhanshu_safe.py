import docx

doc_path = r"c:\Users\LENOVO\Downloads\spesis dataset\review 2\REVIEW-2_AegisSepsis_AI_Sudhanshu_Kumar_Jha.docx"
doc = docx.Document(doc_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().encode('ascii', 'xmlcharrefreplace').decode('ascii')
    if txt:
        print(f"P[{i:03d}] (Style: {p.style.name}): {txt[:120]}")

print("\n=== ALL TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx in range(len(table.rows)):
        row_text = [cell.text.replace('\n', ' ').encode('ascii', 'xmlcharrefreplace').decode('ascii') for cell in table.rows[r_idx].cells]
        print(f"Row {r_idx}: {row_text}")
