import docx

doc_path = r"c:\Users\LENOVO\Downloads\spesis dataset\review 2\REVIEW-2_AegisSepsis_AI_Sudhanshu_Kumar_Jha.docx"
doc = docx.Document(doc_path)
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols")
