import docx

doc_path = r"c:\Users\LENOVO\Downloads\spesis dataset\review 2\REVIEW-2_AegisSepsis_AI_Sudhanshu_Kumar_Jha.docx"
doc = docx.Document(doc_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P[{i}] (Style: {p.style.name}): {p.text[:120]}")

print("\n=== TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx in range(min(5, len(table.rows))):
        row_text = [cell.text.replace('\n', ' ')[:50] for cell in table.rows[r_idx].cells]
        print(f"Row {r_idx}: {row_text}")
