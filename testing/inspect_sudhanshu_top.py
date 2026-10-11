import docx

doc_path = r"c:\Users\LENOVO\Downloads\spesis dataset\review 2\REVIEW-2_AegisSepsis_AI_Sudhanshu_Kumar_Jha.docx"
doc = docx.Document(doc_path)

print("Paragraph count:", len(doc.paragraphs))
for i in range(30):
    print(f"P[{i}]: {repr(doc.paragraphs[i].text)}")
