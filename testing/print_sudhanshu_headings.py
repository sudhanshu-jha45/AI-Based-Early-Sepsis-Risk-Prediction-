import docx

doc_path = r"c:\Users\LENOVO\Downloads\spesis dataset\review 2\REVIEW-2_AegisSepsis_AI_Sudhanshu_Kumar_Jha.docx"
doc = docx.Document(doc_path)

print(f"Total Paragraphs: {len(doc.paragraphs)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().encode('ascii', 'xmlcharrefreplace').decode('ascii')
    if txt and (p.style.name.startswith('Heading') or p.style.name in ['List Paragraph', 'Title'] or txt.isupper() or len(txt) < 80 and ':' in txt):
        print(f"P[{i:03d}] ({p.style.name}): {txt}")
