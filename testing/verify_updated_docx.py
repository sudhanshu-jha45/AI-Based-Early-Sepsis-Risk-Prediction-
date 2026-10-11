import docx

doc = docx.Document(r"c:\Users\LENOVO\Downloads\spesis dataset\REVIEW-2 DOCUMENT FORMAT-Corrected.docx")
print(f"Total Paragraphs in Updated DOCX: {len(doc.paragraphs)}")
print(f"Total Tables in Updated DOCX: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols")
